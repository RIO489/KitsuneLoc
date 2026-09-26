"""JSON-дерево перекладу  <->  робочі книги Excel.

Перекладач бачить лише .xlsx у теці `Переклад`. JSON у `work` — службовий
проміжний шар, який знають тільки скрипти.

Аркуш «Текст»      — усе, що треба перекласти.
Аркуш «Імена»      — імена мовців (Mary Skelter), кожне один раз.
Аркуш «Службове»   — внутрішні ключі гри; їх не чіпати.

Ключ рядка — приховані колонки A (файл) і B (id). Колонки шукаються за
назвою в заголовку, тож порядок/набір видимих колонок можна міняти.
"""
import json, os, re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

import common as locfile
from crystar import tags
from maryskelter import chars as mchars

NAMES = '@names'
COL_SRC, COL_ID, COL_TR, COL_NOTE, COL_HINT = 'файл', 'id', 'Переклад', 'Примітка', 'Підказка'
# фрази, які раніше автоматично писались у «Примітку» — не плутати з нотатками людини
AUTO_NOTES = ('службовий ключ — не перекладати', 'підставлено автоматично — перевір',
              "ім'я мовця — переклад підставиться в усі репліки")
HINTF = Font(color='FF6B7280', size=9)
# спільні об'єкти стилів: створювати їх на кожну клітинку — найповільніше місце
WRAP = Alignment(wrap_text=True, vertical='top')
UNLOCKED = Protection(locked=False)
HID = PatternFill('solid', fgColor='FFF2F2F2')
HEADF = Font(bold=True, color='FFFFFFFF')
HEADFILL = PatternFill('solid', fgColor='FF4A5568')
AUTO = PatternFill('solid', fgColor='FFE8F0FE')

_ID = re.compile(r'^[A-Za-z][A-Za-z0-9_.\-]*$')
_REC = re.compile(r'^\d+:\d+$')            # id репліки в .gbin: запис:зсув поля


COL_STATE, COL_AUTO = 'Стан', 'авто'
ST_TODO, ST_CHECK, ST_DONE = '⬜ не перекладено', '🔵 перевір', '✅ готово'
# Для умовного форматування Excel бере колір з bgColor (end_color), а не fgColor —
# тому задаємо обидва, інакше правило є, а заливки в Excel не видно.
def _cf_fill(rgb):
    return PatternFill(fill_type='solid', start_color=rgb, end_color=rgb)


CHECKFILL = _cf_fill('FFBBD7FF')     # 🔵 перевір — увесь рядок
DONEFILL = _cf_fill('FFCDEFD6')      # ✅ готово — клітинка «Стан»


COL_PIC = 'Як виглядає'
PIC_H = 32                           # мінімальна висота рядка з картинкою, пункти


COL_TERMS = 'Терміни'                # з глосарію (glossary.py): які терміни є в рядку


def columns(game, has_ja=None, has_pic=False, has_terms=False):
    cols = [(COL_SRC, 0), (COL_ID, 0), (COL_STATE, 17), ('Сцена', 16), ('Хто / ключ', 16),
            ('Оригінал (EN)', 55)]
    if has_pic:
        cols.append((COL_PIC, 38))
    if has_ja if has_ja is not None else game == 'crystar':
        cols.append(('Японська', 45))
    cols.append((COL_TR, 60))
    if has_terms:
        cols.append((COL_TERMS, 30))
    cols += [(COL_NOTE, 24), (COL_HINT, 34), (COL_AUTO, 0)]
    return cols


def looks_like_key(s):
    if not s or ' ' in s or '\n' in s or not _ID.match(s):
        return False
    return '_' in s or any(c.isdigit() for c in s) or (len(s) > 3 and s.isupper())


# ---------------------------------------------------------------- читання JSON
def _scene(source):
    parts = source.split('/')
    if source.startswith('Event/') and len(parts) >= 3:     # Event/ev_101010/ev_101010_msg
        return parts[-2]
    base = parts[-1]
    if base.lower() == 'main.cl3' and len(parts) >= 2:      # Neptunia: event/script/0101/main.cl3
        return parts[-2]
    return base.rsplit('.', 1)[0] if '.' in base else base


# службові коди Mary Skelter: printf-формати й #-коди (крім переносу #n)
MSK_CODE = re.compile(r'%[-+ 0#]*\d*(?:\.\d+)?[a-zA-Z]|#[A-Za-mo-z]')


def msk_codes(text):
    import collections
    return collections.Counter(MSK_CODE.findall(text or ''))


def msk_hint(src, cap=None):
    parts = []
    if cap:
        parts.append(f'не більше {cap - 1} байтів (латиниця — 1 Б, кирилиця — 2 Б на літеру)')
    for c in sorted(msk_codes(src)):
        parts.append(f'{c} — підставить гра, лишити' if c.startswith('%') else f'{c} — службовий код, лишити')
    if any('\u3040' <= ch <= '\u9fff' for ch in src or ''):
        parts.append('японський рядок — у англ. версії, ймовірно, не показується')
    return '\n'.join(parts)


# службові коди Neptunia: #FontColor[%u] / #FontColorB, %s %04u %llu %%, <BLANK>
# (перенос рядка в книзі — звичайний Enter, у файлах гри це #n або \n)
NEP_CODE = re.compile(r'#[A-Za-z]+(?:\[[^\]]*\])?|%[-+0#]*\d*(?:\.\d+)?(?:ll|l|h)?[a-zA-Z%]|<[A-Z]+>')


def nep_codes(text):
    import collections
    return collections.Counter(NEP_CODE.findall(text or ''))


def nep_hint(src, cap=None):
    parts = []
    if cap:
        parts.append(f'не більше {cap - 1} символів (поле фіксованої довжини)')
    for c in sorted(nep_codes(src)):
        parts.append(f'{c} — службовий код, лишити')
    return '\n'.join(parts)


def collect(work_dir):
    """[(source, [entry dict])] у стабільному порядку."""
    out = []
    for p in sorted(locfile.walk(work_dir)):
        doc = locfile.load_json(p)
        if doc:
            out.append((doc['source'], doc['entries']))
    return out


# Правила розкладки по книгах: [регулярний вираз по шляху джерела, назва книги].
# Перше правило, що підійшло, визначає книгу. Номер на початку назви — порядок у теці.
# Перевизначити можна файлом `розділи.json` поруч зі скриптами (той самий формат).
DEFAULT_BOOKS = {
    'crystar': [
        [r'^parameter/', '00 Меню та інтерфейс'],
        [r'^Event/ev_10\d{4}/', '01 Пролог'],
        [r'^Event/ev_11\d{4}/', '02 Розділ 1'],
        [r'^Event/ev_12\d{4}/', '03 Розділ 2'],
        [r'^Event/ev_13\d{4}/', '04 Розділ 3'],
        [r'^Event/ev_14\d{4}/', '05 Розділ 4'],
        [r'^Event/ev_15\d{4}/', '06 Розділ 5'],
        [r'^Event/ev_16\d{4}/', '07 Розділ 6'],
        [r'^Event/ev_17\d{4}/', '08 Розділ 7'],
        [r'^Event/ev_18\d{4}/', '09 Фінальний розділ'],
        [r'^Event/ev_2\d{5}/', '10 Повторний прохід (ev_2)'],
        [r'^Event/ev_3\d{5}/', '11 Повторний прохід (ev_3)'],
        [r'^Event/ev_4\d{5}/', '12 Повторний прохід (ev_4)'],
        [r'^Event/ev_calling', '13 Дзвінки'],
        [r'^Event/', '14 Персонажі, NPC, магазин'],
    ],
    'msk': (
        [[r'^TTM1\.bra/Text/', '01 Інтерфейс']]
        # сюжетні сцени 00GGxx: група GG = сотня номера сцени (00…11, 04 і 06 у грі немає)
        + [[rf'/EVENT/DATA/00{g:02d}\d\d\.gbin$', f'{g + 2:02d} Сюжет 00{g:02d}xx'] for g in range(12)]
        + [[r'table\.enc#0$', '17 Спорядження'],
           [r'table\.enc#(1|2|7|44|56)$', '18 Навички та магія'],
           [r'table\.enc#50$', '19 Завдання'],
           [r'table\.enc#(10|25)$', '20 Локації'],
           [r'table\.enc#(3|8|9|11|29|31|41|46|48|58)$', '21 Персонажі та класи'],
           [r'table\.enc#', '22 Інші таблиці'],
           [r'^DLC\.bra/', '23 Описи DLC'],
           [r'^@атлас/', '24 Написи на картинках']]
        + [[r'/EVENT/DATA/10\d{4}\.gbin$', '14 Додаткові сцени 10xxxx'],
           [r'/EVENT/DATA/20\d{4}\.gbin$', '15 Додаткові сцени 20xxxx'],
           [r'/EVENT/DATA/21\d{4}\.gbin$', '16 Короткі репліки 21xxxx']]
    ),
    'nep': (
        [[r'^@атлас/', '22 Написи на картинках'],
         [r'^DLC/', '21 DLC'],
         [r'/database/str\w+\.gstr$', '01 Меню та система'],
         [r'/database/stitem', '02 Предмети'],
         [r'/database/stremake', '03 Створення (Remake)'],
         [r'/database/st(skill|ability)', '04 Навички та вміння'],
         [r'/database/stcharamonster', '05 Монстри'],
         [r'/database/st(charaplayer|avatar|avtmsg)', '06 Персонажі'],
         [r'/database/st(quest|dungeon|lostplace|town|area)', '07 Завдання та локації'],
         [r'/database/', '08 Довідка, галерея, інше']]
        # сюжетні сцени event/script/GGxx: група GG — перші дві цифри номера
        + [[rf'/event/script/{g:02d}\d\d/', f'{10 + g} Сцени {g:02d}xx'] for g in range(10)]
        + [[r'/event/script/', '20 Інші сцени']]
    ),
}
NAMES_BOOK = '00 Імена'
OTHER_BOOK = '99 Інше'
STALE_DIR = '_старі книги'


def book_rules(game):
    p = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'розділи.json')
    if os.path.exists(p):
        try:
            custom = json.load(open(p, encoding='utf-8')).get(game)
            if custom:
                return custom
        except Exception:
            pass
    return DEFAULT_BOOKS.get(game, [])


def plan(game, files):
    """[(назва книги, [(source, entries)])] у порядку назв; порожні книги не створюються."""
    rules = [(re.compile(rx), name) for rx, name in book_rules(game)]
    books = {}
    for f in files:
        name = next((n for rx, n in rules if rx.search(f[0])), OTHER_BOOK)
        books.setdefault(name, []).append(f)
    return sorted(books.items())


def name_ids(entries):
    """id полів, які в .gbin тримають ім'я мовця.

    Ім'я — це ПЕРШЕ рядкове поле файлу (не запису) і лише тоді, коли в записі
    є ще й друге поле з реплікою.
    """
    byrec, offs = {}, set()
    for e in entries:
        eid = e['id']
        if _REC.match(eid):
            rec, fo = eid.split(':')
            byrec.setdefault(rec, {})[fo] = e['src']
            offs.add(int(fo))
    out = {}
    if len(offs) < 2:
        return out, byrec
    first = str(min(offs))
    for rec, fields in byrec.items():
        if len(fields) >= 2 and first in fields:
            out[f'{rec}:{first}'] = fields[first]
    return out, byrec


def _rows(source, entries):
    """[(entry, сцена, контекст, вид)] — вид: 'text' | 'key' | 'name'"""
    scene, ctx, out = _scene(source), '', []
    names, byrec = name_ids(entries)
    for e in entries:
        eid, src = e['id'], e['src']
        if e.get('kind') == 'key' or (e.get('kind') is None and 'ctx' not in e
                                        and looks_like_key(src)):
            ctx = src
            out.append((e, scene, src, 'key'))
            continue
        if eid in names:
            out.append((e, scene, '', 'name'))
            continue
        who = e.get('ctx', ctx)
        if _REC.match(eid):
            fields = byrec.get(eid.split(':')[0], {})
            who = fields[sorted(fields, key=int)[0]] if len(fields) >= 2 else ''
        out.append((e, scene, who, 'text'))
    return out


# ------------------------------------------------------------------ запис XLSX
def _sheet(wb, title, cols, first=False):
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    ws.append([c for c, _w in cols])
    for i, (_c, w) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = w or 12
        if not w:
            ws.column_dimensions[get_column_letter(i)].hidden = True
    for c in ws[1]:
        c.font, c.fill = HEADF, HEADFILL
    ws.freeze_panes = 'D2'
    return ws


def write_book(path, game, book, autofill=True, tagdict=None, names_only=None, has_ja=None,
               terms=None):
    """Одна книга. `names_only` — {ім'я: (переклад, японською)}: тоді це книга «Імена»."""
    has_pic = any(e.get('preview') for _s, es in book for e in es)
    import glossary
    has_terms = bool(terms) and names_only is None
    cols = columns(game, has_ja, has_pic, has_terms)
    with_ja = any(c == 'Японська' for c, _w in cols)
    L = {c: get_column_letter(k + 1) for k, (c, _w) in enumerate(cols)}
    idx = {c: k for k, (c, _w) in enumerate(cols)}
    wb = Workbook()
    ws = _sheet(wb, 'Імена' if names_only is not None else 'Текст', cols, first=True)
    sv = None if names_only is not None else _sheet(wb, 'Службове', cols)

    memory = {}
    if autofill and names_only is None:
        for _s, entries in book:
            for e in entries:
                if e.get('tr'):
                    memory.setdefault(e['src'], e['tr'])

    WRAP_COLS = [idx[c] for c in ('Оригінал (EN)', COL_TR, COL_TERMS, COL_NOTE, COL_HINT) if c in idx]
    if 'Японська' in idx:
        WRAP_COLS.append(idx['Японська'])
    rowno = {}                      # свій лічильник: ws.max_row перебирає всі клітинки

    def add(target, source, e, scene, who, tr, hint, auto=''):
        src, ja = e['src'], e.get('ja', '')
        if game == 'crystar' and tagdict is not None:
            tl = tags.legend(src, ja, tagdict)
            hint = '\n'.join(x for x in (hint, tl) if x)
            ja = tags.resolve(ja, 'ja', tagdict)
        elif game == 'msk':
            hint = '\n'.join(x for x in (hint, e.get('hint'), msk_hint(src, e.get('cap'))) if x)
        elif game == 'nep':
            hint = '\n'.join(x for x in (hint, nep_hint(src, e.get('cap'))) if x)
        r = rowno.get(id(target), 1) + 1
        state = (f'=IF(TRIM({L[COL_TR]}{r})="","{ST_TODO}",'
                 f'IF(AND({L[COL_TR]}{r}={L[COL_AUTO]}{r},{L[COL_NOTE]}{r}=""),"{ST_CHECK}","{ST_DONE}"))')
        row = [source, e['id'], state, scene, who, src]
        if has_pic:
            row.append('')                  # картинку кладемо окремо, поверх клітинки
        if with_ja:
            row.append(ja)
        row.append(tr)
        if has_terms:
            row.append(glossary.hint(src, terms))
        row += [e.get('note', ''), hint, auto]
        target.append(row)
        rowno[id(target)] = r
        # не target[r]: openpyxl тоді щоразу шукає max_column перебором УСІХ
        # клітинок аркуша — на великих книгах це було 3/4 часу кроку «1»
        cells = [target.cell(row=r, column=k + 1) for k in range(len(row))]
        if has_pic and e.get('preview') and os.path.exists(e['preview']):
            from openpyxl.drawing.image import Image as XLImage
            pic = XLImage(e['preview'])
            target.add_image(pic, f'{L[COL_PIC]}{r}')
            target.row_dimensions[r].height = max(PIC_H, pic.height * 0.75 + 4)
        for k in WRAP_COLS:
            cells[k].alignment = WRAP
        for k in range(idx['Сцена'], idx[COL_HINT] + 1):
            cells[k].protection = UNLOCKED
        cells[idx[COL_HINT]].font = HINTF
        return cells

    n_auto = 0
    if names_only is not None:
        for src, (tr, ja) in names_only.items():
            add(ws, NAMES, {'id': src, 'src': src, 'ja': ja}, 'уся гра', '', tr,
                "ім'я мовця — переклад підставиться в усі репліки")
    else:
        for source, entries in book:
            for e, scene, who, kind in _rows(source, entries):
                tr = e.get('tr', '')
                if kind == 'name':
                    continue                       # імена — в окремій книзі «Імена»
                if kind == 'key':
                    cells = add(sv, source, e, scene, who, tr, 'службовий ключ — не перекладати')
                    cells[idx['Оригінал (EN)']].fill = HID
                    continue
                if not tr and e['src'] in memory:
                    tr = memory[e['src']]
                    n_auto += 1
                    add(ws, source, e, scene, who, tr,
                        'підставлено автоматично: такий самий рядок уже перекладено деінде', auto=tr)
                else:
                    add(ws, source, e, scene, who, tr, '', auto=e.get('auto', ''))

    last = get_column_letter(len(cols))
    for sheet in [x for x in (ws, sv) if x is not None]:
        n = sheet.max_row
        if n > 1:
            sheet.auto_filter.ref = f'C1:{L[COL_HINT]}{n}'
            sheet.conditional_formatting.add(
                f'C2:{L[COL_HINT]}{n}',
                FormulaRule(formula=[f'$C2="{ST_CHECK}"'], fill=CHECKFILL, stopIfTrue=True))
            sheet.conditional_formatting.add(
                f'C2:C{n}',
                FormulaRule(formula=[f'$C2="{ST_DONE}"'], fill=DONEFILL, stopIfTrue=True))
        # Захищені лише службові колонки (файл, id, Стан, авто) — решту можна правити
        sheet.protection.sheet = True
        for attr in ('formatCells', 'formatColumns', 'formatRows', 'sort', 'autoFilter',
                     'insertRows', 'insertColumns', 'deleteRows', 'deleteColumns',
                     'insertHyperlinks', 'objects', 'scenarios', 'pivotTables'):
            setattr(sheet.protection, attr, False)

    if sv is not None and not sv.max_row > 1:
        wb.remove(sv)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    return n_auto


def export(work_dir, xlsx_dir, game, autofill=True, progress=None):
    files = collect(work_dir)
    tagdict = tags.build_dict(files) if game == 'crystar' else None
    books = plan(game, files)
    made = []
    has_ja = game == 'crystar' or any('ja' in e for _s, es in files for e in es)
    names = {}
    for source, entries in files:
        for e, _sc, _w, kind in _rows(source, entries):
            if kind != 'name':
                continue
            tr0, ja0 = names.get(e['src'], ('', ''))
            names[e['src']] = (tr0 or e.get('tr', ''), ja0 or e.get('ja', ''))
    skipped = []
    sigs_path = os.path.join(work_dir, BOOK_SIGS)
    try:
        sigs = json.load(open(sigs_path, encoding='utf-8'))
    except (OSError, ValueError):
        sigs = {}
    import glossary
    terms = glossary.load(xlsx_dir)
    base = (_code_sig() + (json.dumps(tagdict, sort_keys=True, ensure_ascii=False) if tagdict else '')
            + json.dumps(terms, sort_keys=True, ensure_ascii=False))
    unchanged = []

    def safe(path, *args, **kw):
        # усе, з чого складається книга, не змінилось — не переписуємо (це
        # десятки секунд на великих книгах і зайвий запис на диск)
        fn = os.path.basename(path)
        sig = _book_sig(base, args, kw)
        if sigs.get(fn, {}).get('sig') == sig and os.path.exists(path):
            unchanged.append(fn)
            return sigs[fn].get('auto', 0)
        # книга відкрита в Excel -> Windows не дає її переписати; пропускаємо лише її
        try:
            n = write_book(path, *args, **kw)
        except PermissionError:
            skipped.append(fn)
            return None
        sigs[fn] = {'sig': sig, 'auto': n}
        return n

    if names:
        path = os.path.join(xlsx_dir, NAMES_BOOK + '.xlsx')
        safe(path, game, [], autofill, tagdict, names_only=dict(sorted(names.items())),
             has_ja=has_ja)
        made.append((path, 0))
    for i, (name, book) in enumerate(books):
        path = os.path.join(xlsx_dir, f'{name}.xlsx')
        # колонка «Японська» лише там, де японський текст справді є (в інтерфейсі MSK його немає)
        book_ja = game == 'crystar' or any('ja' in e for _s, es in book for e in es)
        made.append((path, safe(path, game, book, autofill, tagdict, has_ja=book_ja, terms=terms) or 0))
        if progress:
            progress(i + 1, len(books), name)
    prune_stale(xlsx_dir, [p for p, _n in made])
    try:
        os.makedirs(work_dir, exist_ok=True)
        json.dump(sigs, open(sigs_path, 'w', encoding='utf-8'), ensure_ascii=False, indent=0)
    except OSError:
        pass
    export.skipped = skipped
    export.unchanged = unchanged
    return made


BOOK_SIGS = '_книги.sig'            # у work\<гра>: відбиток даних кожної книги на час запису


def _code_sig():
    """Змінився код, що складає книги, — книги переписуються всі."""
    here = os.path.dirname(os.path.abspath(__file__))
    out = []
    for p in ('sheets.py', os.path.join('crystar', 'tags.py'), os.path.join('maryskelter', 'chars.py')):
        try:
            out.append(str(os.path.getmtime(os.path.join(here, p))))
        except OSError:
            out.append('0')
    return '|'.join(out)


def _book_sig(base, args, kw):
    """md5 усього, з чого write_book складає книгу (разом з картинками-прев'ю)."""
    import hashlib
    h = hashlib.md5(base.encode('utf-8'))
    game, book, autofill = args[0], args[1], args[2]
    h.update(json.dumps([game, autofill, kw.get('has_ja'), kw.get('names_only')],
                        ensure_ascii=False, sort_keys=True).encode('utf-8'))
    for source, entries in book:
        h.update(source.encode('utf-8'))
        h.update(json.dumps(entries, ensure_ascii=False, sort_keys=True).encode('utf-8'))
        for e in entries:
            if e.get('preview') and os.path.exists(e['preview']):
                # вміст, а не час: прев'ю перемальовуються при кожному експорті
                with open(e['preview'], 'rb') as f:
                    h.update(hashlib.md5(f.read()).digest())
    return h.hexdigest()


def prune_stale(xlsx_dir, keep):
    """Книги, яких нова розкладка вже не створює, переносимо в «_старі книги» —
    інакше їхні застарілі рядки змішувались би з новими при читанні."""
    import shutil, time
    keep = {os.path.basename(k) for k in keep}
    stale = [f for f in os.listdir(xlsx_dir)
             if f.endswith('.xlsx') and not f.startswith('~$') and f not in keep]
    if not stale:
        return []
    dst = os.path.join(xlsx_dir, STALE_DIR, time.strftime('%Y-%m-%d_%H%M%S'))
    os.makedirs(dst, exist_ok=True)
    for f in stale:
        shutil.move(os.path.join(xlsx_dir, f), os.path.join(dst, f))
    return stale


# ----------------------------------------------------------------- читання XLSX
BOOK_CACHE_V = 1                    # змінилось те, що читаємо з книги, — +1
BOOK_CACHE = '_книги.cache'         # у work\<гра>: прочитане з книг (за часом зміни й розміром)


def _read_book(path):
    """[(файл, id, переклад, авто | None, примітка | None)] з усіх аркушів книги."""
    out = []
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            it = ws.iter_rows(values_only=True)
            head = next(it, None)
            if not head or COL_SRC not in head or COL_TR not in head:
                continue
            i_s, i_i, i_t = head.index(COL_SRC), head.index(COL_ID), head.index(COL_TR)
            i_n = head.index(COL_NOTE) if COL_NOTE in head else None
            i_a = head.index(COL_AUTO) if COL_AUTO in head else None
            for r in it:
                if not r or not r[i_s]:
                    continue
                tr = r[i_t]
                # Excel — джерело правди: порожня клітинка = перекладу немає
                # (так прибраний у книзі переклад прибирається і з гри)
                tr = tr.replace('\r\n', '\n').strip() if isinstance(tr, str) else ''
                au = (r[i_a] or '') if i_a is not None else None
                n = None
                if i_n is not None:
                    n = r[i_n] if isinstance(r[i_n], str) else ''
                    n = '' if n.strip() in AUTO_NOTES else n.strip()
                out.append((r[i_s], str(r[i_i]), tr, au, n))
    finally:
        wb.close()
    return out


def _book_cache(work_dir):
    import pickle
    try:
        with open(os.path.join(work_dir, BOOK_CACHE), 'rb') as f:
            c = pickle.load(f)
        return c if isinstance(c, dict) else {}
    except Exception:
        return {}


def _save_book_cache(work_dir, cache):
    import pickle
    try:
        os.makedirs(work_dir, exist_ok=True)
        p = os.path.join(work_dir, BOOK_CACHE)
        with open(p + '.tmp', 'wb') as f:
            pickle.dump(cache, f, protocol=pickle.HIGHEST_PROTOCOL)
        os.replace(p + '.tmp', p)
    except OSError:
        pass


def read_into_work(xlsx_dir, work_dir, progress=None):
    """Перенести переклади з усіх .xlsx назад у JSON."""
    trans, names_all, notes, autos, rows = {}, {}, {}, {}, 0
    books = sorted(f for f in os.listdir(xlsx_dir)
                   if f.endswith('.xlsx') and not f.startswith('~$'))
    cache, read_now = _book_cache(work_dir), {}
    for k, fn in enumerate(books):
        path = os.path.join(xlsx_dir, fn)
        st = os.stat(path)
        sig = (BOOK_CACHE_V, st.st_mtime_ns, st.st_size)
        got = cache.get(fn)
        # книгу, що не змінилась з минулого читання, не розбираємо вдруге
        book_rows = got[1] if got and tuple(got[0]) == sig else _read_book(path)
        read_now[fn] = (sig, book_rows)
        for src, eid, tr, au, n in book_rows:
            rows += 1
            if src == NAMES:
                names_all.setdefault(eid, []).append(tr)
            else:
                trans.setdefault(src, {}).setdefault(eid, []).append(tr)
                if au is not None:
                    autos.setdefault(src, {})[eid] = au
                if n is not None:
                    notes.setdefault(src, {})[eid] = n
        if progress:
            progress(k + 1, len(books), fn)
    if read_now != cache:
        _save_book_cache(work_dir, read_now)
    # Одне ім'я може траплятися в кількох книгах (старі книги мали «Імена» в кожній).
    # Заповнене завжди перемагає порожнє; якщо заповнених різних кілька — беремо
    # те, що відрізняється від уже збереженого (тобто свіжу правку).
    docs = []                       # читаємо JSON один раз і тримаємо в пам'яті
    for p in sorted(locfile.walk(work_dir)):
        doc = locfile.load_json(p)
        if doc:
            docs.append((p, doc))
    old_names = {}
    if names_all:
        for _p, doc in docs:
            nids = name_ids(doc['entries'])[0]
            for e in doc['entries']:
                if e['id'] in nids and e.get('tr'):
                    old_names.setdefault(e['src'], e['tr'])
    names = {}
    for src, vals in names_all.items():
        filled = [v for v in vals if v]
        if not filled:
            names[src] = ''
        else:
            fresh = [v for v in filled if v != old_names.get(src)]
            names[src] = fresh[0] if fresh else filled[0]

    def pick(vals, cur):
        filled = [v for v in vals if v]
        if not filled:
            return ''
        fresh = [v for v in filled if v != cur]
        return fresh[0] if fresh else filled[0]

    changed = done = 0
    for p, doc in docs:
        cur_tr = {e['id']: e.get('tr', '') for e in doc['entries']}
        m = {k: pick(v, cur_tr.get(k, '')) for k, v in trans.get(doc['source'], {}).items()}
        nids = name_ids(doc['entries'])[0] if names else {}
        hit = False
        for e in doc['entries']:
            if e['id'] in m:
                t = m[e['id']]
            elif e['id'] in nids and e['src'] in names:
                t = names[e['src']]
            else:
                t = e.get('tr', '')          # рядка немає в жодній книзі — не чіпаємо
            if t != e.get('tr', ''):
                e['tr'], hit = t, True
            au = autos.get(doc['source'], {}).get(e['id'])
            if au is not None:
                # мітка «підставлено автоматично» живе, доки переклад не змінили
                au = au if (au and au == e.get('tr')) else ''
                if au != e.get('auto', ''):
                    if au:
                        e['auto'] = au
                    else:
                        e.pop('auto', None)
                    hit = True
            nt = notes.get(doc['source'], {}).get(e['id'])
            if nt is not None and nt != e.get('note', ''):
                if nt:
                    e['note'] = nt
                else:
                    e.pop('note', None)
                hit = True
            if e.get('tr'):
                done += 1
        if hit:
            json.dump(doc, open(p, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
            changed += 1
    return {'books': len(books), 'rows': rows, 'files': changed, 'translated': done}


# ------------------------------------------------------------------- перевірки
_DIALOG = re.compile(r'/event/script/|/EVENT/DATA/', re.I)


def _width_group(doc, e):
    """Рядки однієї групи показуються в одному й тому ж місці екрана, тож
    найширший оригінал групи — оцінка ширини цього місця. Діалоги — одне
    вікно на всю гру; поля таблиць Neptunia — за зсувом поля в записі; поля
    таблиць Mary Skelter — за місткістю слота (однакові поля — однаковий ліміт)."""
    src = doc['source']
    if _DIALOG.search(src):
        return ('діалог',)
    m = re.match(r'^\d+[.:](\d+)$', e['id'])
    if m:
        return (src, 'поле', m.group(1))
    if e.get('cap'):
        return (src, 'слот', e['cap'])
    return None


def _nlines(s):
    """Кількість рядків; переноси по краях не рахуються — імпорт повертає їх з оригіналу."""
    return s.strip().count('\n') + 1


def _backup_of(work_dir):
    """work\\<гра> -> backup\\<гра> (обидві теки лежать поруч у корені програми)."""
    w = os.path.abspath(work_dir)
    return os.path.join(os.path.dirname(os.path.dirname(w)), 'backup', os.path.basename(w))


def limits(docs, backup_dir, cache_dir):
    """Межі з оригіналів для перевірки ширини: {'wtab', 'game', 'gmax', 'gcount', 'glines'}.
    По групах (див. _width_group): найширший рядок і найбільше рядків в оригіналах."""
    import metrics
    docs = [d for d in docs if d.get('format') != 'atlas']   # написи малюємо самі: ліміти тут не діють
    game = next((d.get('game') for d in docs), None)
    wtab = metrics.table(game, backup_dir, cache_dir) if game else None
    gmax, gcount, glines = {}, {}, {}
    # екрани, де видно репліку: головне вікно й історія діалогів (Neptunia — різні шрифти:
    # advfont і msgfont; MSK — той самий msgfont, але різна ширина). 'id' — ключ DIALOG_WIDTH
    screens = []
    if wtab:
        if game == 'nep':
            adv = metrics.table(game, backup_dir, cache_dir, 'adv')
            if adv:
                screens.append({'назва': 'головне вікно діалогу', 'id': 'adv', 'font': 'adv',
                                'wtab': adv, 'w': []})
            screens.append({'назва': 'вікно історії діалогів', 'id': 'msg', 'font': 'msg',
                            'wtab': wtab, 'w': []})
        else:
            screens.append({'назва': 'вікно діалогу', 'id': 'msg', 'font': 'msg', 'wtab': wtab, 'w': []})
            if (game, 'log') in DIALOG_WIDTH:
                screens.append({'назва': 'вікно історії діалогів', 'id': 'log', 'font': 'msg',
                                'wtab': wtab, 'w': []})
    for doc in docs:
        for e in doc['entries']:
            g = _width_group(doc, e)
            if g is None:
                continue
            gcount[g] = gcount.get(g, 0) + 1
            glines[g] = max(glines.get(g, 0), _nlines(e['src']))
            if wtab:
                w = max(metrics.width(x, wtab, game) for x in e['src'].split('\n'))
                gmax[g] = max(gmax.get(g, 0), w)
                if g == ('діалог',):
                    for sc in screens:
                        sc['w'] += [metrics.width(x, sc['wtab'], game) for x in e['src'].split('\n')]
    for sc in screens:
        # межа вікна — виміряна в грі (DIALOG_WIDTH), інакше не найширший оригінал
        # (там бувають викиди, що й у грі не влазять: «(´・ω・｀) Aaaaaah...» у
        # Neptunia), а 99,9% рядків оригіналу
        w = sorted(sc.pop('w')) or [0]
        sc['lim'] = DIALOG_WIDTH.get((game, sc['id'])) or w[min(len(w) - 1, int(len(w) * DIALOG_PCT))]
    return {'wtab': wtab, 'game': game, 'gmax': gmax, 'gcount': gcount, 'glines': glines,
            'screens': screens}


DIALOG_PCT = 0.999          # частка рядків оригіналу, що мусить уміститися у вікні діалогу
# Ширина вікна діалогу, виміряна в грі (px шрифту гри): тестові рядки, де гра обрізає текст
# (2026-09-26, п'ять рядків на кожне вікно дали ту саму межу ±4 px). Гра не переносить — обрізає.
DIALOG_WIDTH = {('nep', 'adv'): 790,        # головне вікно, advfont: видно 787, обрізано з 794
                ('nep', 'msg'): 726,        # історія діалогів, msgfont: видно 723, обрізано з 730
                # MSK (2026-09-26, рядок-лінійка): головне вікно НЕ обрізає — текст лізе на рамку;
                # внутрішній край вікна ≈ x 1755 з 1920 при масштабі 1,056 екр. px на px шрифту →
                # ≈1290, беремо 1280. Історія (Backlog) обрізає: видно 1032, зникло з 1044
                ('msk', 'msg'): 1280,
                ('msk', 'log'): 1036}


def width_limit(doc, e, ctx, screen=None):
    """(межа ширини в px, скільки рядків уміщає місце | None, ширини рядків оригіналу)
    — за шрифтом гри (ctx['wtab'] є). Для діалогів — у шрифті `screen`
    (ctx['screens'], за замовчуванням перший: головне вікно)."""
    import metrics
    wtab, game = ctx['wtab'], ctx['game']
    g = _width_group(doc, e)
    if g == ('діалог',) and ctx.get('screens'):
        screen = screen or ctx['screens'][0]
        wtab = screen['wtab']
    wsrc = [metrics.width(x, wtab, game) for x in e['src'].split('\n')]
    lines = None
    if g is not None and (g == ('діалог',) or ctx['gcount'].get(g, 0) >= 5):
        lines = ctx['glines'][g]
    if g == ('діалог',):
        lim = screen['lim'] if screen else ctx['gmax'][g]
    elif g is not None and ctx['gcount'].get(g, 0) >= 5:
        lim = max(ctx['gmax'][g], max(wsrc) * 1.1)
    else:
        lim = max(max(wsrc) * 1.35, max(wsrc) + 6 * (wtab.get('n') or 12))
    return lim, lines, wsrc


def check_entry(doc, e, ctx, terms=None, tagdict=None):
    """Попередження до перекладу одного рядка (без source/id); ctx — limits()."""
    import glossary, metrics
    out = []
    tr = e.get('tr')
    if not tr or doc.get('format') == 'atlas':
        return out
    enc = doc.get('encoding')
    is_cry = doc.get('game') == 'crystar'
    wtab, game = ctx['wtab'], ctx['game']
    gmax, gcount, glines = ctx['gmax'], ctx['gcount'], ctx['glines']
    src = e['src']
    for en, ua in glossary.missing(src, tr, terms or []):
        out.append(f'термін «{en}» → «{ua}»: у перекладі не знайдено')
    g = _width_group(doc, e)
    n_src, n_tr = _nlines(src), _nlines(tr)
    if g is not None and (g == ('діалог',) or gcount.get(g, 0) >= 5):
        # переносити можна по-своєму — аби рядків не стало більше, ніж уміщає місце
        if n_tr > glines[g]:
            where = 'вікно діалогу вміщає' if g == ('діалог',) else 'в оригіналах цього поля'
            out.append(f'рядків {n_tr} — {where} не більше {glines[g]}')
    elif n_src != n_tr:
        out.append(f"переносів рядка: було {n_src}, стало {n_tr}")
    if wtab:
        # діалог — кожен екран своїм шрифтом (досить першого, де не влазить)
        for sc in (ctx.get('screens') or [None]) if g == ('діалог',) else [None]:
            lim, _lines, wsrc = width_limit(doc, e, ctx, sc)
            wtr = [metrics.width(x, sc['wtab'] if sc else wtab, game) for x in tr.split('\n')]
            if max(wtr) > lim:
                where = sc['назва'] if sc else 'місце на екрані'
                out.append(f'рядок ширший за {where}: {max(wtr)} px при межі {round(lim)} px '
                           f'(оригінал {max(wsrc)} px) — перенеси рядок або скороти')
                break
    else:
        for a, b in zip(src.split('\n'), tr.split('\n')):
            if len(b) > max(len(a) + 6, len(a) * 1.35):
                out.append(f'рядок довший за оригінал: {len(a)} -> {len(b)} символів')
                break
    if doc.get('game') == 'msk':
        bad = mchars.missing(tr)
        if bad:
            out.append(f"гра не покаже: {' '.join(bad)} (потрібен гліф у шрифті)")
        res = sorted(set(tr) & getattr(mchars, 'RESERVED', set()))
        if res:
            out.append(f"у грі замість {' '.join(res)} з'являться українські "
                       f"літери — прибери ці символи")
    if doc.get('game') == 'nep':
        from neptunia import chars as nchars
        bad = nchars.bad_chars(tr)
        if bad:
            out.append(f"гра не покаже: {' '.join(bad)} (буде «?»)")
        for c, n in (nep_codes(src) - nep_codes(tr)).items():
            out.append(f'бракує коду {c}' + (f' ×{n}' if n > 1 else ''))
        for c in nep_codes(tr) - nep_codes(src):
            out.append(f'зайвий код {c}')
    if e.get('cap') and doc.get('game') == 'nep':
        from neptunia import chars as nchars
        need = len(nchars.encode(tr, 'replace')) + 1
        if need > e['cap']:
            out.append(f"задовго: {need - 1} при ліміті {e['cap'] - 1} символів — "
                       f"лишиться англійським")
    elif e.get('cap'):
        need = len(tr.encode(enc or 'utf-8', 'replace')) + 1
        if need > e['cap']:
            out.append(f"задовго: {need} Б при ліміті {e['cap']} Б")
    if doc.get('game') == 'msk':
        miss = msk_codes(src) - msk_codes(tr)
        for c, n in miss.items():
            out.append(f'бракує коду {c}' + (f' ×{n}' if n > 1 else ''))
        extra = msk_codes(tr) - msk_codes(src)
        for c in extra:
            out.append(f'зайвий код {c} — гра може впасти')
    if is_cry and tagdict is not None:
        for pr in tags.check(src, tr, tagdict):
            out.append(pr)
    if enc == 'cp932':
        bad = sorted({c for c in tr if not _sjis_ok(c)})
        if bad:
            out.append(f"літер немає в кодуванні гри: {' '.join(bad)}")
    return out


def validate(work_dir, backup_dir=None, terms=None):
    """[(source, id, попередження)]; terms — глосарій (glossary.load), щоб ловити
    рядки, де термін в оригіналі є, а його перекладу немає."""
    warn = []
    validate.src = {}                       # (source, id) -> оригінал: для переходу до рядка
    validate.tr = {}                        # (source, id) -> переклад: для «Затверджено»
    files = collect(work_dir)
    tagdict = tags.build_dict(files) if any(s.startswith('parameter/') for s, _ in files) else None
    docs = [d for d in (locfile.load_json(p) for p in sorted(locfile.walk(work_dir))) if d]
    docs = [d for d in docs if d.get('format') != 'atlas']
    ctx = limits(docs, backup_dir or _backup_of(work_dir), os.path.dirname(os.path.abspath(work_dir)))
    for doc in docs:
        for e in doc['entries']:
            if not e.get('tr'):
                continue
            validate.src[(doc['source'], e['id'])] = e['src']
            validate.tr[(doc['source'], e['id'])] = e['tr']
            for m in check_entry(doc, e, ctx, terms, tagdict):
                warn.append((doc['source'], e['id'], m))
    return warn


# ------------------------------------------------------------ «Затверджено»
APPROVED = '_затверджено.json'      # у теці книг гри: рішення перекладача, як і самі книги


def approval_key(source, eid, msg):
    """Ключ попередження без чисел (ширина в px тощо може трохи змінитись)."""
    return f'{source}\t{eid}\t' + re.sub(r'\d+(?:\.\d+)?', '#', msg)


def load_approved(xlsx_dir):
    """{ключ попередження: переклад, для якого його затвердили}."""
    try:
        with open(os.path.join(xlsx_dir, APPROVED), encoding='utf-8') as f:
            d = json.load(f)
        return d if isinstance(d, dict) else {}
    except (OSError, ValueError):
        return {}


def save_approved(xlsx_dir, approved):
    os.makedirs(xlsx_dir, exist_ok=True)
    p = os.path.join(xlsx_dir, APPROVED)
    with open(p + '.tmp', 'w', encoding='utf-8') as f:
        json.dump(approved, f, ensure_ascii=False, indent=1)
    os.replace(p + '.tmp', p)


def split_approved(warn, xlsx_dir):
    """(показати, затверджені) — затвердження діє, доки переклад рядка той самий."""
    ok = load_approved(xlsx_dir)
    tr = getattr(validate, 'tr', {})
    shown, hidden = [], []
    for w in warn:
        s, i, m = w
        k = approval_key(s, i, m)
        (hidden if k in ok and ok[k] == tr.get((s, i)) else shown).append(w)
    return shown, hidden


def _sjis_ok(c):
    try:
        c.encode('cp932')
        return True
    except UnicodeEncodeError:
        return False


def progress(xlsx_dir, work_dir=None):
    """Скільки перекладено — по книгах і разом, прямо з .xlsx.

    Повертає {'books': [...], 'done', 'total', 'auto', 'prev'}, де auto —
    рядки, підставлені автоматично й ще не перевірені людиною.
    """
    books, done, total, auto = [], 0, 0, 0
    if not os.path.isdir(xlsx_dir):
        return {'books': [], 'done': 0, 'total': 0, 'auto': 0, 'prev': None}
    for fn in sorted(f for f in os.listdir(xlsx_dir)
                     if f.endswith('.xlsx') and not f.startswith('~$')):
        b_done, b_total, b_auto = book_stats(os.path.join(xlsx_dir, fn))
        if b_total:
            books.append({'book': fn[:-5], 'done': b_done,
                          'total': b_total, 'auto': b_auto})
            done += b_done; total += b_total; auto += b_auto
    res = {'books': books, 'done': done, 'total': total, 'auto': auto, 'prev': None}
    if work_dir:
        res['prev'] = _snapshot(work_dir, done, total)
    return res


def book_stats(path):
    """(перекладено, усього, підставлено автоматично й не перевірено) в одній книзі."""
    b_done = b_total = b_auto = 0
    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        head = next(it, None)
        if not head or COL_SRC not in head or COL_TR not in head:
            continue
        i_s, i_t = head.index(COL_SRC), head.index(COL_TR)
        i_n = head.index(COL_NOTE) if COL_NOTE in head else None
        i_a = head.index(COL_AUTO) if COL_AUTO in head else None
        for r in it:
            if not r or not r[i_s]:
                continue
            b_total += 1
            tr = r[i_t].strip() if isinstance(r[i_t], str) else ''
            if not tr:
                continue
            b_done += 1
            a = r[i_a] if i_a is not None and isinstance(r[i_a], str) else ''
            n = r[i_n] if i_n is not None and isinstance(r[i_n], str) else ''
            if a and tr == a.strip() and not n.strip():
                b_auto += 1
    wb.close()
    return b_done, b_total, b_auto


# ------------------------------------------------------ перехід до рядка книги
def locate(xlsx_dir, game, source, eid, src=None):
    """Де в книгах лежить рядок: (шлях книги, аркуш, номер рядка) або None.
    Спершу — книга за правилами розкладки, далі книга «Імена» (ім'я мовця
    там один раз, ключ — сам англійський текст), далі всі інші."""
    if not os.path.isdir(xlsx_dir):
        return None
    rules = [(re.compile(rx), name) for rx, name in book_rules(game)]
    first = next((n for rx, n in rules if rx.search(source)), OTHER_BOOK) + '.xlsx'
    books = [f for f in sorted(os.listdir(xlsx_dir))
             if f.endswith('.xlsx') and not f.startswith('~$')]
    order = ([first] if first in books else []) + \
            ([NAMES_BOOK + '.xlsx'] if NAMES_BOOK + '.xlsx' in books and src else []) + \
            [f for f in books if f != first and f != NAMES_BOOK + '.xlsx']
    for fn in order:
        want = (NAMES, src) if fn == NAMES_BOOK + '.xlsx' else (source, eid)
        wb = load_workbook(os.path.join(xlsx_dir, fn), read_only=True)
        try:
            for ws in wb.worksheets:
                it = ws.iter_rows(values_only=True)
                head = next(it, None)
                if not head or COL_SRC not in head or COL_ID not in head:
                    continue
                i_f, i_i = head.index(COL_SRC), head.index(COL_ID)
                for k, r in enumerate(it, 2):
                    if r and r[i_f] == want[0] and str(r[i_i]) == str(want[1]):
                        col = get_column_letter(head.index(COL_TR) + 1) if COL_TR in head else 'A'
                        return os.path.join(xlsx_dir, fn), ws.title, f'{col}{k}'
        finally:
            wb.close()
    return None


def _xml_attr(tag, name, value):
    """Поставити атрибут у відкривальний тег XML (рядком — без перезапису
    простору імен, який Excel не любить)."""
    if value is None:
        return re.sub(rf'\s{name}="[^"]*"', '', tag)
    if re.search(rf'\s{name}="', tag):
        return re.sub(rf'(\s{name}=")[^"]*"', rf'\g<1>{value}"', tag)
    end = '/>' if tag.endswith('/>') else '>'
    return tag[:-len(end)] + f' {name}="{value}"' + end


def goto_cell(path, sheet, cell):
    """Зробити так, щоб книга відкрилась на аркуші `sheet` з виділеною
    клітинкою `cell`: правимо лише вигляд (workbook.xml і sheetView аркуша)
    прямо в архіві .xlsx — openpyxl при перезаписі губить картинки."""
    import zipfile
    m = re.match(r'([A-Z]+)(\d+)$', cell)
    col, row = m.group(1), int(m.group(2))
    with zipfile.ZipFile(path) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    parts = {i.filename: d for i, d in items}
    wbx = parts['xl/workbook.xml'].decode('utf-8')
    rels = parts['xl/_rels/workbook.xml.rels'].decode('utf-8')
    from xml.sax.saxutils import escape
    tabs = re.findall(r'<(?:\w+:)?sheet\b[^>]*>', wbx)
    idx = next(k for k, t in enumerate(tabs) if f'name="{escape(sheet)}"' in t)
    targets = []
    for t in tabs:
        rid = re.search(r'\br:id="([^"]+)"|\bid="([^"]+)"', t)
        rid = rid.group(1) or rid.group(2)
        rel = re.search(rf'<Relationship\b[^>]*\bId="{rid}"[^>]*>', rels).group(0)
        tg = re.search(r'Target="([^"]+)"', rel).group(1)
        targets.append(tg.lstrip('/') if tg.startswith('/') else 'xl/' + tg)
    wbx = re.sub(r'<(?:\w+:)?workbookView\b[^>]*>',
                 lambda mm: _xml_attr(_xml_attr(mm.group(0), 'activeTab', idx), 'firstSheet', None),
                 wbx, count=1)
    parts['xl/workbook.xml'] = wbx.encode('utf-8')
    for k, name in enumerate(targets):
        x = parts[name].decode('utf-8')
        x = re.sub(r'<(?:\w+:)?sheetView\b[^>]*>',
                   lambda mm: _xml_attr(mm.group(0), 'tabSelected', 1 if k == idx else None), x)
        if k == idx:
            pane = re.search(r'<(?:\w+:)?pane\b[^>]*/>', x)
            ysplit = int(float(re.search(r'ySplit="([\d.]+)"', pane.group(0)).group(1))) \
                if pane and 'ySplit=' in pane.group(0) else 0
            top = max(ysplit + 1, row - 3)
            where = f'{col}{row}'
            x = re.sub(r'<(?:\w+:)?selection\b[^>]*/>', '', x)
            if pane:
                active = re.search(r'activePane="([^"]+)"', pane.group(0))
                active = active.group(1) if active else 'bottomRight'
                xs = re.search(r'xSplit="([\d.]+)"', pane.group(0))
                left = get_column_letter(int(float(xs.group(1))) + 1) if xs else 'A'
                new = _xml_attr(pane.group(0), 'topLeftCell', f'{left}{top}')
                x = x.replace(pane.group(0), new +
                              f'<selection pane="{active}" activeCell="{where}" sqref="{where}"/>', 1)
            else:
                sv = re.search(r'<sheetView\b[^>]*>', x).group(0)
                new = _xml_attr(sv, 'topLeftCell', f'A{top}')
                sel = f'<selection activeCell="{where}" sqref="{where}"/>'
                new = new[:-2] + '>' + sel + '</sheetView>' if new.endswith('/>') else new + sel
                x = x.replace(sv, new, 1)
        parts[name] = x.encode('utf-8')
    tmp = path + '.new'
    with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as z:
        for info, _d in items:
            z.writestr(info, parts[info.filename])
    os.replace(tmp, path)


def _snapshot(work_dir, done, total):
    """Запамʼятати сьогоднішній результат і повернути попередній."""
    import datetime
    path = os.path.join(work_dir, 'прогрес.json')
    hist = []
    if os.path.exists(path):
        try:
            hist = json.load(open(path, encoding='utf-8'))
        except Exception:
            hist = []
    prev = hist[-1] if hist else None
    now = {'коли': datetime.datetime.now().strftime('%Y-%m-%d %H:%M'),
           'перекладено': done, 'усього': total}
    if not prev or prev.get('перекладено') != done:
        hist.append(now)
        hist = hist[-200:]
        try:
            os.makedirs(work_dir, exist_ok=True)
            json.dump(hist, open(path, 'w', encoding='utf-8'),
                      ensure_ascii=False, indent=1)
        except Exception:
            pass
    return prev


def bar(done, total, width=24):
    n = round(width * done / total) if total else 0
    return '\u2588' * n + '\u2591' * (width - n)


def report(pr):
    """Текстовий звіт про прогрес — рівними колонками, для вікна програми."""
    if not pr['total']:
        return ['Книг ще немає — спершу натисни «1. Дістати текст з гри».']
    pc = 100 * pr['done'] / pr['total']
    out = [f"Перекладено {pr['done']:,} з {pr['total']:,} рядків — {pc:.1f}%"
           .replace(',', ' '),
           bar(pr['done'], pr['total'], 40)]
    if pr['prev']:
        d = pr['done'] - pr['prev'].get('перекладено', 0)
        if d > 0:
            out.append(f"+{d} рядків з {pr['prev']['коли']}")
    if pr['auto']:
        out.append(f"з них підставлено автоматично й ще не перевірено: {pr['auto']}")
    out.append('')
    w = max(len(b['book']) for b in pr['books'])
    for b in pr['books']:
        p = 100 * b['done'] / b['total']
        mark = ' ' if b['done'] < b['total'] else '\u2713'
        out.append(f"{mark} {b['book']:<{w}}  {bar(b['done'], b['total'], 16)} "
                   f"{p:5.1f}%  {b['done']}/{b['total']}")
    left = pr['total'] - pr['done']
    out.append('')
    out.append(f'Лишилось {left} рядків.' if left else 'Усе перекладено. Красень.')
    return out
