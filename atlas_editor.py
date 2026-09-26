# -*- coding: utf-8 -*-
"""Вікно «Написи на картинках»: переклад з живим прев'ю.

Пишеш переклад — поруч одразу видно, як кнопка виглядатиме в грі (усі її
варіанти: звичайна, вибрана тощо). Зберігається в ту саму книгу Excel
(Mary Skelter — «24 Написи на картинках», Neptunia — «22 Написи на
картинках»), тож Excel лишається джерелом правди, а кнопка «2. Залити
переклад у гру» підхопить усе як завжди.
"""
import os, queue, threading
import tkinter as tk
from tkinter import ttk, messagebox

from PIL import Image, ImageTk

from maryskelter import atlas as atl, dds
from maryskelter.bra import Bra
from maryskelter.cl3 import Cl3

BOOKS = {'msk': '24 Написи на картинках.xlsx', 'nep': '22 Написи на картинках.xlsx'}
SHEET = 'Текст'
BG = (32, 32, 40, 255)
MAX_W = 380             # ширина прев'ю в пікселях екрана
MAX_VARIANTS = 4        # скільки різних варіантів кнопки показувати
DELAY = 180             # мс тиші після набору, перш ніж перемальовувати


class Atlases:
    """Розкодовані атласи з чистих оригіналів (backup, інакше тека гри)."""

    def __init__(self, backup_dir, game_dir):
        self.dirs = [d for d in (backup_dir, game_dir) if d]
        self.cache = {}
        self.lock = threading.Lock()

    def get(self, src, mark):
        with self.lock:
            if src not in self.cache:
                arc, _, name = src.partition('/')
                path = next((os.path.join(d, arc) for d in self.dirs
                             if os.path.exists(os.path.join(d, arc))), None)
                if path is None:
                    raise FileNotFoundError(arc)
                cl3 = Cl3(Bra.read_some(path, [name])[name])
                stem = mark.get('текстура')
                img = dds.decode(bytes(atl.pair(cl3, stem)[1][1]))
                self.cache[src] = (img, atl.frames(cl3, stem))
            return self.cache[src]


class NepTextures:
    """Neptunia: текстури .tid з архівів .pac (backup, інакше тека гри);
    рамка кадру завжди явна в розмітці, тож нарізка порожня."""

    def __init__(self, backup_dir, game_dir):
        self.dirs = [d for d in (backup_dir, game_dir) if d]
        self.cache, self.pacs = {}, {}
        self.lock = threading.Lock()

    def get(self, src, mark):
        from neptunia import atlas as natl
        from neptunia.pac import Pac
        from neptunia.tid import Tid
        with self.lock:
            if src not in self.cache:
                arc, inner = natl.split_src(src)
                if arc not in self.pacs:
                    path = next((os.path.join(d, *arc.split('/')) for d in self.dirs
                                 if os.path.exists(os.path.join(d, *arc.split('/')))), None)
                    if path is None:
                        raise FileNotFoundError(arc)
                    self.pacs[arc] = Pac(path)
                self.cache[src] = (Tid(self.pacs[arc].read(inner)).image(), {})
            return self.cache[src]


def variants(marks, key):
    """[(джерело, кадр, spec)] для ключа — без повторів за стилем і розміром."""
    out, seen = [], set()
    for src, mark in marks.items():
        for i, spec in sorted(mark['кадри'].items(), key=atl.order):
            if atl.key_of(spec) != key:
                continue
            sig = (spec['стиль'], tuple(spec.get('рамка') or ()), src if spec.get('рамка') else '')
            if sig in seen:
                continue
            seen.add(sig)
            out.append((src, i, spec))
    return out[:MAX_VARIANTS]


def tuned(spec, edit=None, real=None):
    """Кадр з пробними змінами: підібраний шрифт (`edit` — правки стилю) і
    «справжні літери» (`real`: True/False, None — як у розмітці)."""
    spec = dict(spec)
    if edit:
        spec['правки'] = dict(spec.get('правки', {}), **edit)
    if real is not None:
        if real:
            spec['літери з оригіналу'] = True
        else:
            spec.pop('літери з оригіналу', None)
    return spec


def render(atlases, marks, styles, key, text, game='msk', edit=None, real=None):
    """[(оригінал, результат, масштаб, [попередження])] для кожного варіанта."""
    res = []
    for src, i, spec in variants(marks, key):
        spec = tuned(spec, edit, real)
        if game == 'nep' and text:
            from neptunia import atlas as natl
            text_i = natl.text_for(spec, text)   # розрізані слова: «!!» окремим спрайтом
        else:
            text_i = text
        img, boxes = atlases.get(src, marks[src])
        box = atl.box_of(i, spec, boxes)
        orig = img.crop(box)
        if text:
            # малюємо на копії всього атласу: тло "шаблон" береться з іншого кадру
            work = img.copy()
            info = {}
            warn = atl.draw(work, box, spec, styles, text_i, info)
            res.append((orig, work.crop(box), info.get('масштаб', 1.0), warn))
        else:
            res.append((orig, orig, 1.0, []))
    return res


def to_photo(im):
    bg = Image.new('RGBA', im.size, BG)
    bg.alpha_composite(im)
    if bg.width > MAX_W:
        bg = bg.resize((MAX_W, max(1, round(bg.height * MAX_W / bg.width))), Image.LANCZOS)
    return ImageTk.PhotoImage(bg.convert('RGB'))


class Editor(tk.Toplevel):
    def __init__(self, app):
        super().__init__(app)
        self.app = app
        work, xl, _out, bk = app.dirs()
        self.game = app.cur['game']
        self.book_name = BOOKS[self.game]
        self.book = os.path.join(xl, self.book_name)
        if not os.path.exists(self.book):
            self.destroy()
            raise RuntimeError(f'Книги «{self.book_name[:-5]}» ще немає — '
                               'спершу натисни «1. Дістати текст з гри».')
        if self.game == 'nep':
            from neptunia import atlas as natl
            self.marks = natl.load_marks()
            self.atlases = NepTextures(bk, app.root_dir())
        else:
            self.marks = {k: v for k, v in atl.load_json('написи.json').items()
                          if not k.startswith('_')}
            self.atlases = Atlases(bk, app.root_dir())
        self.styles = atl.load_json('стилі.json')
        self.rows = self._read_book()           # [{id, src, tr, where}]
        self.edits = {}                          # id -> новий переклад (ще не збережений)
        self.cur = None
        self.gen = 0                             # номер рендеру: застарілі результати відкидаємо
        self.pending = None
        self.photos = []

        self.title('Написи на картинках — переклад з прев\'ю')
        self.geometry('1180x720')
        self.minsize(900, 520)
        c = app_colors(app)
        self.configure(bg=c['bg'])
        self._build(c)
        self._fill()
        self.protocol('WM_DELETE_WINDOW', self._close)
        if self.rows:
            first = self.tree.get_children()[0]
            self.tree.selection_set(first)
            self.tree.focus(first)

    # ---------------------------------------------------------------- книга
    def _read_book(self):
        from openpyxl import load_workbook
        wb = load_workbook(self.book, read_only=True, data_only=True)
        ws = wb[SHEET]
        it = ws.iter_rows(values_only=True)
        head = list(next(it))
        ix = {h: head.index(h) for h in ('id', 'Оригінал (EN)', 'Переклад', 'Хто / ключ')}
        rows = []
        for r in it:
            if not r or r[ix['id']] is None:
                continue
            rows.append({'id': str(r[ix['id']]), 'src': r[ix['Оригінал (EN)']] or '',
                         'tr': (r[ix['Переклад']] or '').strip() if isinstance(r[ix['Переклад']], str) else '',
                         'where': r[ix['Хто / ключ']] or ''})
        wb.close()
        return rows

    def _save(self):
        if not self.edits:
            return True
        lock = os.path.join(os.path.dirname(self.book), '~$' + self.book_name)
        if os.path.exists(lock):
            messagebox.showwarning('Книга відкрита в Excel',
                                   f'Закрий «{self.book_name[:-5]}» в Excel і збережи ще раз — '
                                   'інакше Excel потім перезапише ці зміни.', parent=self)
            return False
        from openpyxl import load_workbook
        wb = load_workbook(self.book)
        ws = wb[SHEET]
        head = [c.value for c in ws[1]]
        i_id, i_tr = head.index('id'), head.index('Переклад')
        n = 0
        for row in ws.iter_rows(min_row=2):
            k = str(row[i_id].value)
            if k in self.edits:
                row[i_tr].value = self.edits[k] or None
                n += 1
        try:
            wb.save(self.book)
        except PermissionError:
            messagebox.showwarning('Не вдалося зберегти',
                                   'Книгу зайнято (відкрита в Excel?). Закрий її й спробуй ще раз.',
                                   parent=self)
            return False
        for r in self.rows:
            if r['id'] in self.edits:
                r['tr'] = self.edits[r['id']]
        self.edits.clear()
        self.saved.set(f'Збережено в книгу: {n}. Далі — «2. Залити переклад у гру».')
        self._fill(keep=True)
        return True

    def _close(self):
        if self.edits:
            ans = messagebox.askyesnocancel('Незбережені зміни',
                                            f'Змінено написів: {len(self.edits)}. Зберегти в книгу?',
                                            parent=self)
            if ans is None:
                return
            if ans and not self._save():
                return
        self.destroy()

    # ---------------------------------------------------------------- вигляд
    def _build(self, c):
        st = ttk.Style()
        st.configure('Pics.Treeview', background=c.get('panel', '#ffffff'),
                     fieldbackground=c.get('panel', '#ffffff'), foreground=c.get('fg', '#1a1a1a'),
                     rowheight=22)
        st.map('Pics.Treeview', background=[('selected', c.get('accent', '#1f4f82'))],
               foreground=[('selected', '#ffffff')])
        top = ttk.Frame(self)
        top.pack(fill='x', padx=10, pady=(10, 4))
        ttk.Label(top, text='Пошук:').pack(side='left')
        self.q = tk.StringVar()
        self.q.trace_add('write', lambda *_: self._fill())
        ttk.Entry(top, textvariable=self.q, width=30).pack(side='left', padx=6)
        self.only_todo = tk.BooleanVar(value=False)
        ttk.Checkbutton(top, text='лише неперекладені', variable=self.only_todo,
                        command=self._fill).pack(side='left', padx=6)
        self.count = tk.StringVar()
        ttk.Label(top, textvariable=self.count).pack(side='right')

        body = ttk.Panedwindow(self, orient='horizontal')
        body.pack(fill='both', expand=True, padx=10, pady=4)

        left = ttk.Frame(body)
        self.tree = ttk.Treeview(left, columns=('src', 'tr'), show='headings', selectmode='browse',
                                 style='Pics.Treeview')
        self.tree.heading('src', text='Оригінал')
        self.tree.heading('tr', text='Переклад')
        self.tree.column('src', width=160)
        self.tree.column('tr', width=160)
        sb = ttk.Scrollbar(left, orient='vertical', command=self.tree.yview)
        self.tree.configure(yscrollcommand=sb.set)
        self.tree.pack(side='left', fill='both', expand=True)
        sb.pack(side='left', fill='y')
        self.tree.bind('<<TreeviewSelect>>', lambda e: self._select())
        self.tree.tag_configure('todo', foreground=c['dim'])
        self.tree.tag_configure('edit', foreground=c['warn'])
        body.add(left, weight=0)

        right = ttk.Frame(body, padding=(12, 0, 0, 0))
        body.add(right, weight=1)
        self.head = tk.StringVar()
        ttk.Label(right, textvariable=self.head, font=('Segoe UI', 12, 'bold')).pack(anchor='w')
        self.where = tk.StringVar()
        ttk.Label(right, textvariable=self.where, style='Hint.TLabel').pack(anchor='w', pady=(0, 6))

        er = ttk.Frame(right)
        er.pack(fill='x')
        ttk.Label(er, text='Переклад:').pack(side='left')
        self.text = tk.StringVar()
        self.entry = ttk.Entry(er, textvariable=self.text, font=('Segoe UI', 12))
        self.entry.pack(side='left', fill='x', expand=True, padx=6)
        self.text.trace_add('write', lambda *_: self._typed())
        self.entry.bind('<Return>', lambda e: self._next())
        self.entry.bind('<Down>', lambda e: self._next())
        self.entry.bind('<Up>', lambda e: self._next(-1))
        ttk.Button(er, text='Лишити оригінал', command=lambda: self.text.set('')).pack(side='left')

        # вигляд напису: підбір шрифту за оригіналом і «справжні літери»
        tools = ttk.Frame(right)
        tools.pack(fill='x', pady=(6, 0))
        self.real = tk.BooleanVar(value=False)
        ttk.Checkbutton(tools, text='Справжні літери з оригіналу',
                        variable=self.real, command=self._real_toggled).pack(side='left')
        ttk.Button(tools, text='Глянути різні шрифти…', command=self._gallery).pack(side='left', padx=8)
        ttk.Button(tools, text='Повернути стандартний', command=self._standard).pack(side='left')
        self.b_apply = ttk.Button(tools, text='Записати в розмітку', command=self._apply_look)
        self.b_apply.pack(side='left', padx=8)
        self.b_apply.state(['disabled'])
        self.look = tk.StringVar()
        ttk.Label(right, textvariable=self.look, style='Hint.TLabel').pack(anchor='w')
        self.suggest = {}                        # ключ -> підібрані правки стилю (ще не в розмітці)

        self.note = tk.StringVar()
        self.note_lbl = ttk.Label(right, textvariable=self.note)
        self.note_lbl.pack(anchor='w', pady=(4, 4))
        self.warn_color, self.ok_color = c['warn'], c['dim']

        grid = ttk.Frame(right)
        grid.pack(fill='both', expand=True)
        ttk.Label(grid, text='Оригінал', style='Hint.TLabel').grid(row=0, column=0, sticky='w')
        ttk.Label(grid, text='У грі буде', style='Hint.TLabel').grid(row=0, column=1, sticky='w',
                                                                   padx=(12, 0))
        self.cells = []
        for k in range(MAX_VARIANTS):
            a = tk.Label(grid, bg=c['bg'], bd=0)
            b = tk.Label(grid, bg=c['bg'], bd=0)
            a.grid(row=k + 1, column=0, sticky='nw', pady=4)
            b.grid(row=k + 1, column=1, sticky='nw', pady=4, padx=(12, 0))
            self.cells.append((a, b))

        bot = ttk.Frame(self)
        bot.pack(fill='x', padx=10, pady=(4, 10))
        self.saved = tk.StringVar(value='Enter / ↓ — наступний напис, ↑ — попередній.')
        ttk.Label(bot, textvariable=self.saved, style='Hint.TLabel').pack(side='left')
        ttk.Button(bot, text='Закрити', command=self._close).pack(side='right')
        ttk.Button(bot, text='Зберегти в книгу', command=self._save).pack(side='right', padx=8)

    # ---------------------------------------------------------------- список
    def _value(self, r):
        return self.edits.get(r['id'], r['tr'])

    def _fill(self, keep=False):
        q = self.q.get().strip().lower()
        sel = self.cur
        self.tree.delete(*self.tree.get_children())
        shown = 0
        for r in self.rows:
            v = self._value(r)
            if self.only_todo.get() and v:
                continue
            if q and q not in r['src'].lower() and q not in v.lower():
                continue
            tag = 'edit' if r['id'] in self.edits else ('' if v else 'todo')
            self.tree.insert('', 'end', iid=r['id'], values=(r['src'], v or '—'), tags=(tag,))
            shown += 1
        done = sum(1 for r in self.rows if self._value(r))
        self.count.set(f'перекладено {done} з {len(self.rows)}')
        if sel and self.tree.exists(sel):
            self.tree.selection_set(sel)
            self.tree.see(sel)

    def _select(self):
        s = self.tree.selection()
        if not s or s[0] == self.cur:
            return
        self.cur = s[0]
        r = next(x for x in self.rows if x['id'] == self.cur)
        self.head.set(r['src'])
        self.where.set(f'Де: {r["where"]}')
        self._loading = True
        self.text.set(self._value(r))
        self.real.set(any(s.get('літери з оригіналу') for _src, _i, s in variants(self.marks, self.cur)))
        self._loading = False
        self._look_status()
        self.entry.focus_set()
        self.entry.icursor('end')
        self._schedule(0)

    def _next(self, step=1):
        items = self.tree.get_children()
        if not items:
            return 'break'
        k = items.index(self.cur) if self.cur in items else -1
        nxt = items[max(0, min(len(items) - 1, k + step))]
        self.tree.selection_set(nxt)
        self.tree.see(nxt)
        return 'break'

    # ---------------------------------------------------------------- прев'ю
    def _typed(self):
        if getattr(self, '_loading', False) or not self.cur:
            return
        r = next(x for x in self.rows if x['id'] == self.cur)
        v = self.text.get().strip()
        if v == r['tr']:
            self.edits.pop(r['id'], None)
        else:
            self.edits[r['id']] = v
        if self.tree.exists(r['id']):
            tag = 'edit' if r['id'] in self.edits else ('' if v else 'todo')
            self.tree.item(r['id'], values=(r['src'], v or '—'), tags=(tag,))
        self._schedule(DELAY)

    def _schedule(self, ms):
        if self.pending:
            self.after_cancel(self.pending)
        self.pending = self.after(ms, self._render)

    def _render(self):
        self.pending = None
        self.gen += 1
        gen, key, text = self.gen, self.cur, self.text.get().strip()
        self.note.set('малюю…')

        edit, real = self.suggest.get(key), self.real.get()

        def work():
            try:
                res = render(self.atlases, self.marks, self.styles, key, text, self.game, edit, real)
                err = None
            except Exception as ex:                                   # noqa: BLE001
                res, err = [], str(ex)
            self.after(0, lambda: self._show(gen, text, res, err))

        threading.Thread(target=work, daemon=True).start()

    def _show(self, gen, text, res, err):
        if gen != self.gen or not self.winfo_exists():
            return
        self.photos = []
        for k, (a, b) in enumerate(self.cells):
            if k < len(res):
                o, n = to_photo(res[k][0]), to_photo(res[k][1])
                self.photos += [o, n]
                a.configure(image=o)
                b.configure(image=n)
                a.grid()
                b.grid()
            else:
                a.configure(image='')
                b.configure(image='')
                a.grid_remove()                 # порожня мітка малюється сірою рисочкою
                b.grid_remove()
        if err:
            self.note.set(f'Не вдалося намалювати: {err}')
            self.note_lbl.configure(foreground=self.warn_color)
            return
        if not text:
            self.note.set('Порожньо — у грі лишиться оригінальний напис.')
            self.note_lbl.configure(foreground=self.ok_color)
            return
        scale = min((x[2] for x in res), default=1.0)
        warns = [w for x in res for w in x[3]]
        if warns:
            self.note.set(f'Задовге: напис зменшено до {scale:.0%} — спробуй коротше.')
            self.note_lbl.configure(foreground=self.warn_color)
        elif scale < 0.97:
            how = 'трохи стиснуто' if scale >= 0.85 else 'помітно зменшено'
            self.note.set(f'Вміщається, {how} ({scale:.0%} від оригінального розміру).')
            self.note_lbl.configure(foreground=self.ok_color)
        else:
            self.note.set('Вміщається без зменшення.')
            self.note_lbl.configure(foreground=self.ok_color)


    # ------------------------------------------------------- вигляд напису
    def _marked_real(self, key):
        return any(s.get('літери з оригіналу') for _src, _i, s in variants(self.marks, key))

    def _look_status(self):
        key = self.cur
        changed = key in self.suggest or self.real.get() != self._marked_real(key)
        self.b_apply.state(['!disabled'] if changed else ['disabled'])
        if key in self.suggest:
            e = self.suggest[key]
            var = e.get('варіація') or {}
            self.look.set(f'Пробний шрифт: {e["шрифт"]}' +
                          (f', товщина {var["wght"]}' if 'wght' in var else '') +
                          f', нахил {e["нахил"]}, розтяг {e["розтяг"]}, розрядка {e.get("розрядка", 0)}'
                          ' — ще не записано.')
        elif changed:
            self.look.set('«Справжні літери» змінено лише для прев\'ю — ще не записано.')
        else:
            self.look.set('')

    def _real_toggled(self):
        self._look_status()
        self._schedule(0)

    def _gallery(self):
        if self.cur and variants(self.marks, self.cur):
            FontGallery(self, self.cur)

    def _chosen(self, key, edit):
        """Шрифт, вибраний у галереї: лише для прев'ю, поки не записано."""
        if edit is None:
            self.suggest.pop(key, None)
        else:
            self.suggest[key] = edit
        if key == self.cur:
            self._look_status()
            self._schedule(0)

    @staticmethod
    def _restored(spec):
        """Кадр з правками стилю, що були до запису шрифту (None — шрифт не записували)."""
        if 'правки до шрифту' not in spec:
            return None
        spec = dict(spec)
        old = spec.pop('правки до шрифту')
        if old:
            spec['правки'] = old
        else:
            spec.pop('правки', None)
        return spec

    def _standard(self):
        """Прибрати пробний шрифт, а якщо шрифт уже записано в розмітку —
        повернути там правки стилю, що були до нього."""
        key = self.cur
        self.suggest.pop(key, None)
        mod, marks = self._markup()
        n = 0
        for src, mark in marks.items():
            if src.startswith('_'):
                continue
            for i, spec in mark['кадри'].items():
                new = self._restored(spec) if atl.key_of(spec) == key else None
                if new is not None:
                    mark['кадри'][i] = new
                    n += 1
        if n:
            mod.save_marks(marks)
            self.marks = {k: v for k, v in marks.items() if not k.startswith('_')}
        self._look_status()
        self.look.set(f'Повернуто стандартний шрифт у розмітці: кадрів {n}.' if n else
                      'Стандартний шрифт (у розмітці інший і не записували).')
        self._schedule(0)

    def _apply_all(self, edit):
        """Один шрифт для ВСІХ написів гри (edit — лише шрифт і варіація; нахил і
        ширина кожного стилю лишаються свої). edit None — усім стандартний."""
        mod, marks = self._markup()
        n = 0
        for src, mark in marks.items():
            if src.startswith('_'):
                continue
            for i, spec in mark['кадри'].items():
                if edit is None:
                    new = self._restored(spec)
                    if new is None:
                        continue
                else:
                    if 'правки до шрифту' not in spec:
                        spec = dict(spec, **{'правки до шрифту': spec.get('правки')})
                    new = tuned(spec, edit)
                mark['кадри'][i] = new
                n += 1
        mod.save_marks(marks)
        if edit:
            from maryskelter import fontlib
            fontlib.adopt(edit['шрифт'])
        self.suggest.clear()
        self.marks = {k: v for k, v in marks.items() if not k.startswith('_')}
        self._look_status()
        self.look.set((f'Шрифт {edit["шрифт"]} записано для всіх написів' if edit else
                       'Повернуто стандартні шрифти всім написам') + f': кадрів {n}.')
        self._schedule(0)

    def _markup(self):
        """(модуль інструмента розмітки, уся розмітка з диска)."""
        import importlib.util
        tool = 'розмітка_неп.py' if self.game == 'nep' else 'розмітка.py'
        sp = importlib.util.spec_from_file_location('_tool', os.path.join(atl.DIR, tool))
        mod = importlib.util.module_from_spec(sp)
        sp.loader.exec_module(mod)
        return mod, mod.load_marks()

    def _apply_look(self):
        """Записати вибраний шрифт і «справжні літери» в розмітку для всіх
        кадрів цього напису (атлас/написи.json чи атлас/нептун.json). Попередні
        правки стилю зберігаються в "правки до шрифту" — для «Повернути стандартний»."""
        key = self.cur
        mod, marks = self._markup()
        edit, real, n = self.suggest.get(key), self.real.get(), 0
        for src, mark in marks.items():
            if src.startswith('_'):
                continue
            for i, spec in mark['кадри'].items():
                if atl.key_of(spec) != key:
                    continue
                if edit and 'правки до шрифту' not in spec:
                    spec = dict(spec, **{'правки до шрифту': spec.get('правки')})
                mark['кадри'][i] = tuned(spec, edit, real)
                n += 1
        mod.save_marks(marks)
        if edit:
            from maryskelter import fontlib
            fontlib.adopt(edit['шрифт'])
        self.suggest.pop(key, None)
        self.marks = {k: v for k, v in marks.items() if not k.startswith('_')}
        self._look_status()
        self.look.set(f'Записано в розмітку: кадрів {n}.')
        self._schedule(0)


def app_colors(app):
    import importlib
    mod = importlib.import_module('__main__')
    themes = getattr(mod, 'THEMES', None) or {'light': {'bg': '#f5f5f5', 'dim': '#6a6a6a',
                                                        'warn': '#8a6100'}}
    return themes.get(getattr(app, 'theme', 'light'), next(iter(themes.values())))


class FontGallery(tk.Toplevel):
    """«Глянути різні шрифти»: той самий напис кожним шрифтом бібліотеки —
    стандартний першим. Товщину, нахил і ширину можна крутити для всіх разом;
    клік — вибрати, «Взяти вибраний» (або подвійний клік) — у прев'ю редактора."""

    COLS = 2
    CARD_W = 420                            # ширина картинки в картці, пікселі екрана

    def __init__(self, editor, key):
        super().__init__(editor)
        from maryskelter import fontlib
        self.ed, self.key = editor, key
        self.src, self.i, self.spec = variants(editor.marks, key)[0]
        st = dict(editor.styles[self.spec['стиль']])
        st.update(self.spec.get('правки', {}))
        st.update(editor.suggest.get(key) or {})
        var = st.get('варіація')
        self.fonts = [(None, None)] + fontlib.font_files()     # None — стандартний
        self.cards, self.photos, self.sel = [], {}, 0
        self.gen, self.pending = 0, None
        self.q = queue.Queue()

        c = app_colors(editor.app)
        self.c = c
        self.title(f'Різні шрифти — {key}')
        self.geometry('1000x760')
        self.configure(bg=c['bg'])

        top = ttk.Frame(self, padding=(10, 10, 10, 4))
        top.pack(fill='x')
        img, boxes = editor.atlases.get(self.src, editor.marks[self.src])
        orig = img.crop(atl.box_of(self.i, self.spec, boxes))
        # дрібні написи збільшуємо (до 2×), великі — зменшуємо до ширини картки
        self.zoom = min(2.0, self.CARD_W / max(1, orig.width))
        self.orig_photo = self._photo(orig)
        ttk.Label(top, text='Оригінал:').grid(row=0, column=0, sticky='nw')
        tk.Label(top, image=self.orig_photo, bg=c['bg'], bd=0).grid(row=0, column=1, columnspan=6,
                                                                     sticky='w', padx=8)
        ttk.Label(top, text='Текст:').grid(row=1, column=0, sticky='w', pady=(8, 0))
        self.text = tk.StringVar(value=editor.text.get().strip() or self.spec['текст'])
        ttk.Entry(top, textvariable=self.text, width=28).grid(row=1, column=1, sticky='w',
                                                              padx=8, pady=(8, 0))
        self.weight = tk.DoubleVar(value=var.get('wght', 700) if isinstance(var, dict) else 700)
        self.slant = tk.DoubleVar(value=st.get('нахил', 0))
        self.stretch = tk.DoubleVar(value=st.get('розтяг', 1.0))
        self.labels = {}
        for col, (name, v, lo, hi) in enumerate((('Товщина', self.weight, 100, 900),
                                                 ('Нахил', self.slant, -0.1, 0.4),
                                                 ('Ширина', self.stretch, 0.7, 1.4)), start=2):
            box = ttk.Frame(top)
            box.grid(row=1, column=col, padx=8, pady=(8, 0), sticky='w')
            self.labels[name] = ttk.Label(box, width=14)
            self.labels[name].pack(anchor='w')
            ttk.Scale(box, from_=lo, to=hi, variable=v, length=150,
                      command=lambda _v: self._changed()).pack()
        self.text.trace_add('write', lambda *_: self._changed())
        ttk.Label(top, text='Повзунки діють на всі шрифти, крім «Стандартного». Товщина — лише '
                            'для шрифтів з кількома товщинами; у …-Bold, …-Black вона зашита у файлі.',
                  style='Hint.TLabel').grid(row=2, column=0, columnspan=7, sticky='w', pady=(4, 0))

        body = ttk.Frame(self)
        body.pack(fill='both', expand=True, padx=10)
        self.canvas = tk.Canvas(body, bg=c['bg'], highlightthickness=0)
        sb = ttk.Scrollbar(body, orient='vertical', command=self.canvas.yview)
        self.canvas.configure(yscrollcommand=sb.set)
        sb.pack(side='right', fill='y')
        self.canvas.pack(side='left', fill='both', expand=True)
        self.inner = tk.Frame(self.canvas, bg=c['bg'])
        self.canvas.create_window((0, 0), window=self.inner, anchor='nw')
        self.inner.bind('<Configure>',
                        lambda e: self.canvas.configure(scrollregion=self.canvas.bbox('all')))
        self.bind('<MouseWheel>', self._wheel)
        for k, (_d, fn) in enumerate(self.fonts):
            card = tk.Frame(self.inner, bg=c['bg'], bd=0, highlightthickness=2,
                            highlightbackground=c['bg'], padx=6, pady=4)
            card.grid(row=k // self.COLS, column=k % self.COLS, sticky='nw', padx=4, pady=4)
            name = tk.Label(card, text='Стандартний (як зараз)' if fn is None else fn.rsplit('.', 1)[0],
                            bg=c['bg'], fg=c.get('fg', '#1a1a1a'), anchor='w',
                            font=('Segoe UI', 9, 'bold' if fn is None else 'normal'))
            name.pack(anchor='w')
            pic = tk.Label(card, bg=c['bg'], fg=c.get('dim', '#6a6a6a'), bd=0, text='…')
            pic.pack(anchor='w')
            for w in (card, name, pic):
                w.bind('<Button-1>', lambda e, n=k: self._select(n))
                w.bind('<Double-Button-1>', lambda e, n=k: (self._select(n), self._take()))
                w.bind('<MouseWheel>', self._wheel)
            self.cards.append((card, pic))

        bot = ttk.Frame(self, padding=(10, 6, 10, 10))
        bot.pack(fill='x')
        self.status = tk.StringVar()
        ttk.Label(bot, textvariable=self.status, style='Hint.TLabel').pack(side='left')
        ttk.Button(bot, text='Закрити', command=self._close).pack(side='right')
        ttk.Button(bot, text='Взяти вибраний', command=self._take).pack(side='right', padx=8)
        ttk.Button(bot, text='Для всіх написів…', command=self._take_all).pack(side='right')
        self.protocol('WM_DELETE_WINDOW', self._close)
        self._select(0)
        self._changed(0)
        self.after(40, self._poll)

    # ------------------------------------------------------------------
    def _wheel(self, e):
        self.canvas.yview_scroll(-1 if e.delta > 0 else 1, 'units')
        return 'break'

    def _select(self, n):
        c = self.c
        self.cards[self.sel][0].configure(highlightbackground=c['bg'])
        self.sel = n
        self.cards[n][0].configure(highlightbackground=c.get('accent', '#1f4f82'))
        fn = self.fonts[n][1]
        self.status.set('Вибрано: ' + ('стандартний шрифт' if fn is None else fn) +
                        '. «Взяти вибраний» (або подвійний клік) — показати в редакторі.')

    def _edit(self, d, fn):
        """Правки стилю для шрифту з бібліотеки при поточних повзунках."""
        from maryskelter import fontlib
        if fn is None:
            return None
        fontlib.register(d, fn)
        return {'шрифт': fn, 'варіація': fontlib.variation(os.path.join(d, fn), self.weight.get()),
                'нахил': round(self.slant.get(), 3), 'розтяг': round(self.stretch.get(), 3)}

    def _changed(self, delay=350):
        self.labels['Товщина'].configure(text=f'Товщина {self.weight.get():.0f}')
        self.labels['Нахил'].configure(text=f'Нахил {self.slant.get():.2f}')
        self.labels['Ширина'].configure(text=f'Ширина {self.stretch.get():.2f}')
        if self.pending:
            self.after_cancel(self.pending)
        self.pending = self.after(delay, self._start)

    def _start(self):
        """Перемалювати всі картки у фоні (попередній прохід зупиняється сам)."""
        self.pending = None
        self.gen += 1
        gen, ed = self.gen, self.ed
        text = self.text.get().strip() or self.spec['текст']
        edits = [self._edit(d, fn) for d, fn in self.fonts]
        spec0, real, game = self.spec, ed.real.get(), ed.game

        def work():
            img, boxes = ed.atlases.get(self.src, ed.marks[self.src])
            box = atl.box_of(self.i, spec0, boxes)
            canvas = img.copy()                 # "шаблон" бере тло з іншого місця атласу
            clean = img.crop(box)
            t = text
            if game == 'nep':
                from neptunia import atlas as natl
                t = natl.text_for(spec0, text)
            for k, edit in enumerate(edits):
                if gen != self.gen:
                    return
                canvas.paste(clean, box[:2])
                try:
                    atl.draw(canvas, box, tuned(spec0, edit, real), ed.styles, t)
                    self.q.put((gen, k, canvas.crop(box), None))
                except Exception as ex:                               # noqa: BLE001
                    self.q.put((gen, k, None, str(ex)))
            self.q.put((gen, None, None, None))

        self.status.set('Малюю…')
        threading.Thread(target=work, daemon=True).start()

    def _poll(self):
        if not self.winfo_exists():
            return
        try:
            while True:
                gen, k, im, err = self.q.get_nowait()
                if gen != self.gen:
                    continue
                if k is None:
                    self._select(self.sel)
                    continue
                pic = self.cards[k][1]
                if im is None:
                    pic.configure(image='', text=f'не вдалося: {err}', fg=self.c.get('warn', '#8a6100'))
                    continue
                self.photos[k] = self._photo(im)
                pic.configure(image=self.photos[k], text='')
        except queue.Empty:
            pass
        self.after(40, self._poll)

    def _photo(self, im):
        bg = Image.new('RGBA', im.size, BG)
        bg.alpha_composite(im)
        if self.zoom != 1.0:
            bg = bg.resize((max(1, round(bg.width * self.zoom)), max(1, round(bg.height * self.zoom))),
                           Image.LANCZOS)
        return ImageTk.PhotoImage(bg.convert('RGB'))

    def _take_all(self):
        """Вибраний шрифт — одразу в розмітку всіх написів гри."""
        d, fn = self.fonts[self.sel]
        edit = self._edit(d, fn)
        if edit is None:
            q = ('Повернути стандартні шрифти всім написам, яким шрифт записували?')
        else:
            var = edit.get('варіація') or {}
            q = (f'Записати шрифт {fn}' + (f' (товщина {var["wght"]})' if 'wght' in var else '') +
                 ' для ВСІХ написів цієї гри?\n\nНахил і ширина кожного стилю лишаться свої. '
                 'Скасувати можна тут же: картка «Стандартний» → «Для всіх написів…».')
            edit = {k: edit[k] for k in ('шрифт', 'варіація')}
        if not messagebox.askyesno('Для всіх написів', q, parent=self):
            return
        self.ed._apply_all(edit)
        self._close()

    def _take(self):
        d, fn = self.fonts[self.sel]
        self.ed._chosen(self.key, self._edit(d, fn))
        self._close()

    def _close(self):
        self.gen += 1                           # зупинити фоновий прохід
        self.destroy()
